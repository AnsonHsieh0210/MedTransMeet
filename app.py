import io
import re
import time
from pathlib import Path
import streamlit as st
from google import genai
from google.genai import types
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

# 1. 網頁頁面設定
st.set_page_config(page_title="醫療專科 雲端 AI 批次語音轉錄系統", layout="centered")
st.title("🩺 醫療專科 雲端 AI 批次語音轉錄系統")

MODEL = "gemini-3.8-flash"

# 依副檔名判斷 MIME 類型（瀏覽器回報的 type 不可靠，例如 m4a 常為 audio/x-m4a 或空值）
MIME_TYPES = {
    ".m4a": "audio/mp4",
    ".mp3": "audio/mpeg",
    ".wav": "audio/wav",
    ".mp4": "video/mp4",
}

# 檔案在 Gemini 雲端處理的最長等待時間（秒）
PROCESSING_TIMEOUT = 600

# Word 檔使用的中文字型
DOCX_FONT = "Microsoft JhengHei"

# 文章類型與對應的撰寫提示詞
ARTICLE_STYLES = {
    "會議紀錄": (
        "請根據以下會議逐字稿，撰寫一份正式的繁體中文（台灣用語）會議紀錄，包含：\n"
        "1. 會議主題\n2. 會議摘要（約 200 字）\n3. 討論重點（條列）\n"
        "4. 決議事項\n5. 待辦事項（含負責人與期限，逐字稿未提及者標示「未指定」）\n"
        "請忠於逐字稿內容，不要捏造未提及的資訊。醫療資訊專有名詞（如 FHIR、HL7、LOINC）請保持正確拼寫與大寫。"
    ),
    "新聞稿": (
        "請根據以下會議逐字稿，撰寫一篇繁體中文（台灣用語）新聞稿，包含吸引人的標題、導言（5W1H）、"
        "內文三到五段，以及結語。語氣專業、客觀，適合對外發布。\n"
        "請忠於逐字稿內容，不要捏造未提及的資訊或引述。醫療資訊專有名詞請保持正確拼寫與大寫。"
    ),
    "專題文章": (
        "請根據以下會議逐字稿，撰寫一篇結構完整的繁體中文（台灣用語）專題文章，包含標題、前言、"
        "數個附小標的段落，以及結論。適合刊登於院內刊物或部落格，讓非與會者也能理解討論內容與意義。\n"
        "請忠於逐字稿內容，不要捏造未提及的資訊。醫療資訊專有名詞請保持正確拼寫與大寫，首次出現時可簡短說明。"
    ),
}


def _add_markdown(doc, text):
    """將模型輸出的簡易 Markdown（標題、條列、粗體）轉為 Word 段落"""
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if heading:
            doc.add_heading(heading.group(2).replace("**", ""), level=min(len(heading.group(1)) + 1, 4))
            continue
        if re.match(r"^[-*•]\s+", stripped):
            para = doc.add_paragraph(style="List Bullet")
            stripped = re.sub(r"^[-*•]\s+", "", stripped)
        elif re.match(r"^\d+[.、)]\s*", stripped):
            para = doc.add_paragraph(style="List Number")
            stripped = re.sub(r"^\d+[.、)]\s*", "", stripped)
        else:
            para = doc.add_paragraph()
        # **粗體** 轉為粗體文字
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", stripped)):
            if part:
                para.add_run(part).bold = i % 2 == 1


def build_docx(results):
    """產生 Word 檔：results 為 [(檔名, {"transcript": ..., "article": ...}), ...]"""
    doc = Document()
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        style = doc.styles[style_name]
        style.font.name = DOCX_FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
    doc.styles["Normal"].font.size = Pt(12)

    for index, (filename, result) in enumerate(results):
        if index > 0:
            doc.add_page_break()
        doc.add_heading(Path(filename).stem, level=0)
        if result.get("article"):
            doc.add_heading("文章", level=1)
            _add_markdown(doc, result["article"])
        doc.add_heading("逐字稿原文", level=1)
        for line in result["transcript"].splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_txt(filename, result):
    text = f"=== 檔案名稱: {filename} ===\n"
    if result.get("article"):
        text += f"\n【文章】\n{result['article']}\n"
    text += f"\n【逐字稿原文】\n{result['transcript']}\n"
    return text


# 2. 由使用者自行填入 Google AI API Key（僅存在本次瀏覽工作階段，不會儲存，也不使用網站管理者的金鑰）
with st.sidebar:
    st.markdown("### 🔑 Gemini API Key")
    api_key = st.text_input(
        "請輸入您自己的 Gemini API Key",
        type="password",
        help="金鑰僅用於本次瀏覽工作階段，關閉頁面後即消失，不會被儲存。",
    )
    st.markdown("[👉 免費申請 API Key（Google AI Studio）](https://aistudio.google.com/apikey)")
    st.caption(f"使用模型：`{MODEL}`。轉錄費用由您的 API Key 帳號支付。")

if not api_key:
    st.warning("🔑 請先在左側欄輸入您自己的 Gemini API Key 才能開始使用。")
    st.markdown(
        "**如何取得 API Key：**\n"
        "1. 前往 [Google AI Studio](https://aistudio.google.com/apikey) 並以 Google 帳號登入\n"
        "2. 點選「Create API key」建立金鑰\n"
        "3. 複製金鑰並貼到左側欄\n\n"
        "⚠️ 若錄音含病患個資，建議使用已啟用付費方案的金鑰（免費層資料可能被用於改善模型）。"
    )
else:
    client = genai.Client(api_key=api_key.strip())

    # 💡 關鍵更動：開啟 accept_multiple_files=True
    uploaded_files = st.file_uploader(
        "請選擇並上傳一個或多個會議音訊檔案",
        type=["m4a", "mp3", "wav", "mp4"],
        accept_multiple_files=True
    )

    # 用來存放所有檔案轉錄結果的字典：{檔名: {"transcript": 逐字稿, "article": 文章}}
    if "batch_outputs" not in st.session_state:
        st.session_state["batch_outputs"] = {}
    # 雲端刪除失敗的檔案（需提醒使用者手動處理）
    if "delete_failures" not in st.session_state:
        st.session_state["delete_failures"] = []

    if uploaded_files:
        st.write(f"📋 已選取 {len(uploaded_files)} 個檔案，準備進行批次處理。")

        # 3. 建立專屬的醫療術語強化提示詞
        with st.expander("🩺 醫療資訊專家提示詞設定 (已啟用)", expanded=False):
            medical_prompt = st.text_area(
                "給 AI 的指導原則：",
                value=(
                    "你是一位精通智慧醫療、醫療資訊系統（HIS/PACS）與國際醫療資料標準的專業記錄員。\n"
                    "請將這段音訊精準轉錄為繁體中文（台灣習慣用語）。\n"
                    "【特別注意】音訊中包含大量醫療資訊與標準化編碼的專業術語，請務必正確拼寫，例如：\n"
                    "- 醫療資訊交換與編碼標準：FHIR, HL7, DICOM, SNOMED CT, RxNorm, LOINC, ICD-10\n"
                    "- 檢驗與影像系統：PACS (Picture Archiving and Communication System), RIS, HIS, LIS\n"
                    "若遇到英文縮寫，請保持大寫（如 LOINC、FHIR）。不確定的語句請結合醫療與檢驗數據上下文進行合理推斷，保持專門術語的精確度。"
                ),
                height=200
            )

        # 文章撰寫設定
        write_article = st.checkbox("✍️ 轉錄後自動撰寫文章", value=True)
        if write_article:
            article_style = st.selectbox("文章類型", list(ARTICLE_STYLES))
            with st.expander("📝 文章撰寫提示詞設定", expanded=False):
                article_prompt = st.text_area(
                    "給 AI 的撰寫指示：",
                    value=ARTICLE_STYLES[article_style],
                    height=200,
                    key=f"article_prompt_{article_style}",
                )

        # 4. 開始批次執行轉錄
        if st.button("🚀 開始批次智慧醫療轉錄", type="primary"):
            # 建立一個進度條
            progress_bar = st.progress(0)
            status_text = st.empty()

            # 清空舊的紀錄
            st.session_state["batch_outputs"] = {}
            st.session_state["delete_failures"] = []

            for index, uploaded_file in enumerate(uploaded_files):
                filename = uploaded_file.name
                # 同名檔案加上編號，避免結果互相覆蓋
                result_key = filename
                dup = 2
                while result_key in st.session_state["batch_outputs"]:
                    result_key = f"{Path(filename).stem} ({dup}){Path(filename).suffix}"
                    dup += 1
                step = f"({index+1}/{len(uploaded_files)}): **{result_key}**"
                status_text.markdown(f"⏳ 正在上傳 {step}")

                gemini_file = None
                result = {"transcript": "", "article": ""}
                try:
                    suffix = Path(filename).suffix.lower()
                    mime_type = MIME_TYPES.get(suffix)
                    if mime_type is None:
                        raise ValueError(f"不支援的檔案格式「{suffix or '無副檔名'}」，僅支援 {', '.join(MIME_TYPES)}")

                    # 透過 Files API 上傳（避開 inline 資料約 20MB 的請求上限，支援長時間會議錄音）
                    uploaded_file.seek(0)
                    gemini_file = client.files.upload(
                        file=uploaded_file,
                        config=types.UploadFileConfig(mime_type=mime_type, display_name=filename),
                    )

                    # 等待雲端處理完成（設有逾時，避免單一檔案卡住整個批次）
                    status_text.markdown(f"⏳ 雲端處理中 {step}")
                    deadline = time.monotonic() + PROCESSING_TIMEOUT
                    while gemini_file.state == types.FileState.PROCESSING:
                        if time.monotonic() > deadline:
                            raise TimeoutError(f"雲端處理超過 {PROCESSING_TIMEOUT // 60} 分鐘仍未完成")
                        time.sleep(2)
                        gemini_file = client.files.get(name=gemini_file.name)
                    if gemini_file.state == types.FileState.FAILED:
                        raise RuntimeError("Gemini 檔案處理失敗")

                    # 呼叫 Gemini 3.8 Flash 模型
                    status_text.markdown(f"⏳ 轉錄中 {step}")
                    response = client.models.generate_content(
                        model=MODEL,
                        contents=[medical_prompt, gemini_file]
                    )

                    # response.text 在內容被安全機制阻擋或無輸出時為 None
                    if not response.text:
                        raise RuntimeError("模型未回傳轉錄內容（可能遭安全機制阻擋或輸出為空）")
                    result["transcript"] = response.text

                    # 依逐字稿撰寫文章（失敗時保留逐字稿，只在文章欄位標示錯誤）
                    if write_article:
                        status_text.markdown(f"✍️ 撰寫文章中 {step}")
                        try:
                            article_response = client.models.generate_content(
                                model=MODEL,
                                contents=[article_prompt, f"【會議逐字稿】\n{result['transcript']}"]
                            )
                            if not article_response.text:
                                raise RuntimeError("模型未回傳文章內容")
                            result["article"] = article_response.text
                        except Exception as e:
                            result["article"] = f"❌ 文章撰寫發生錯誤: {e}"

                except Exception as e:
                    result["transcript"] = f"❌ 轉錄發生錯誤: {e}"

                finally:
                    # 轉錄後立即刪除雲端檔案（醫療音訊不留存，否則預設保留 48 小時）
                    if gemini_file is not None:
                        try:
                            client.files.delete(name=gemini_file.name)
                        except Exception as e:
                            st.session_state["delete_failures"].append(f"{result_key}（{gemini_file.name}）：{e}")

                # 將結果存入 session_state
                st.session_state["batch_outputs"][result_key] = result

                # 更新進度條
                progress_bar.progress((index + 1) / len(uploaded_files))

            status_text.success("✨ 所有檔案批次處理完成！")
            progress_bar.empty()

        # 5. 顯示批次結果與下載區塊
        if st.session_state["batch_outputs"]:
            st.write("---")
            st.markdown("### 📥 轉錄結果下載與預覽")

            if st.session_state["delete_failures"]:
                st.warning(
                    "⚠️ 以下音訊未能從 Gemini 雲端刪除（預設將保留 48 小時），請至 Google AI Studio 手動刪除：\n\n"
                    + "\n".join(f"- {f}" for f in st.session_state["delete_failures"])
                )

            # 建立打包合併的文字
            combined_text = ""

            # 用分頁 (Tabs) 的方式呈現各個檔案的結果，介面比較整齊
            outputs = st.session_state["batch_outputs"]
            file_names = list(outputs.keys())
            tabs = st.tabs(file_names)

            for i, filename in enumerate(file_names):
                with tabs[i]:
                    result = outputs[filename]
                    if result["article"]:
                        st.text_area(f"{filename} 的文章", value=result["article"], height=300, key=f"art_{filename}")
                    st.text_area(f"{filename} 的逐字稿原文", value=result["transcript"], height=300, key=f"txt_{filename}")

                    # 個別下載按鈕（自動對齊檔名）
                    stem_name = Path(filename).stem
                    col_docx, col_txt = st.columns(2)
                    with col_docx:
                        st.download_button(
                            label=f"📄 下載 Word：{stem_name}.docx",
                            data=build_docx([(filename, result)]),
                            file_name=f"{stem_name}.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key=f"docx_{filename}"
                        )
                    with col_txt:
                        st.download_button(
                            label=f"📥 下載文字檔：{stem_name}.txt",
                            data=build_txt(filename, result),
                            file_name=f"{stem_name}.txt",
                            mime="text/plain",
                            key=f"dl_{filename}"
                        )

                    # 併入總文本
                    combined_text += build_txt(filename, result) + "\n"

            # 提供一鍵下載全部合併包的功能
            st.write("---")
            col_docx, col_txt = st.columns(2)
            with col_docx:
                st.download_button(
                    label="📦 下載全部合併 Word (combined_all.docx)",
                    data=build_docx(list(outputs.items())),
                    file_name="combined_all.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary"
                )
            with col_txt:
                st.download_button(
                    label="📦 下載全部合併文字檔 (combined_all.txt)",
                    data=combined_text,
                    file_name="combined_all.txt",
                    mime="text/plain",
                    type="secondary"
                )
